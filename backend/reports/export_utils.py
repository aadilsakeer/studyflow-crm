import csv
import io

from openpyxl import Workbook


def export_reports_csv(reports):
    buffer = io.StringIO()
    writer = csv.writer(buffer)

    sections = (
        ('Lead Source ROI', reports['lead_source_roi'], [
            'source', 'leads', 'converted', 'conversion_rate', 'revenue',
        ]),
        ('Telecaller Performance', reports['telecaller_performance'], [
            'name', 'leads', 'calls', 'converted', 'conversion_rate',
        ]),
        ('Counsellor Performance', reports['counsellor_performance'], [
            'name', 'leads', 'converted', 'conversion_rate',
        ]),
        ('Country Performance', reports['country_performance'], [
            'country', 'students',
        ]),
        ('University Performance', reports['university_performance'], [
            'university', 'applications',
        ]),
    )

    for title, rows, fields in sections:
        writer.writerow([title])
        writer.writerow(fields)

        for row in rows:
            writer.writerow([row.get(field, '') for field in fields])

        writer.writerow([])

    writer.writerow(['Offer Conversion'])
    for key, value in reports['offer_conversion'].items():
        writer.writerow([key, value])

    writer.writerow([])
    writer.writerow(['Visa Success'])
    for key, value in reports['visa_success'].items():
        writer.writerow([key, value])

    writer.writerow([])
    writer.writerow(['Revenue'])
    for key, value in reports['revenue'].items():
        writer.writerow([key, value])

    return buffer.getvalue()


def export_reports_xlsx(reports):
    workbook = Workbook()
    workbook.remove(workbook.active)

    sheet_map = (
        ('Lead Source ROI', reports['lead_source_roi'], [
            'source', 'leads', 'converted', 'conversion_rate', 'revenue',
        ]),
        ('Telecaller Performance', reports['telecaller_performance'], [
            'name', 'leads', 'calls', 'converted', 'conversion_rate',
        ]),
        ('Counsellor Performance', reports['counsellor_performance'], [
            'name', 'leads', 'converted', 'conversion_rate',
        ]),
        ('Country Performance', reports['country_performance'], [
            'country', 'students',
        ]),
        ('University Performance', reports['university_performance'], [
            'university', 'applications',
        ]),
    )

    for title, rows, fields in sheet_map:
        sheet = workbook.create_sheet(title[:31])
        sheet.append(fields)

        for row in rows:
            sheet.append([row.get(field, '') for field in fields])

    summary = workbook.create_sheet('Summary')
    summary.append(['Metric', 'Value'])

    for key, value in reports['offer_conversion'].items():
        summary.append([f'offer_{key}', value])

    for key, value in reports['visa_success'].items():
        summary.append([f'visa_{key}', value])

    for key, value in reports['revenue'].items():
        summary.append([f'revenue_{key}', value])

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
